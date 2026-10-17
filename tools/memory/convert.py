#!/usr/bin/env python3
"""Package the memory workbook without rewriting any source text."""

import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


HEADERS = ("STT", "Chương", "Phần", "Chủ đề / Nội dung", "Nội dung chi tiết", "Mẹo nhớ")
ROOT = Path(__file__).resolve().parents[2]


def source_records(path, sheet_name):
    raw = path.read_bytes()
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Không tìm thấy sheet {sheet_name!r}.")
        sheet = workbook[sheet_name]
        if tuple(sheet.cell(2, column).value for column in range(1, 7)) != HEADERS:
            raise ValueError("Hàng 2 phải có đúng 6 tiêu đề: " + ", ".join(HEADERS))
        title = sheet.cell(1, 1).value
        if not isinstance(title, str) or not title.strip():
            raise ValueError("Ô A1 thiếu tiêu đề.")
        records, seen_stt = [], set()
        for cells in sheet.iter_rows(min_row=3):
            if all(cell.value is None for cell in cells):
                continue
            source_row = cells[0].row
            if any(cell.value is not None for cell in cells[6:]):
                raise ValueError(f"Hàng {source_row}: có nội dung ngoài 6 cột quy định.")
            if any(cell.data_type in {"f", "e"} for cell in cells[:6]):
                raise ValueError(f"Hàng {source_row}: không chấp nhận công thức hoặc ô lỗi.")
            values = tuple(cell.value for cell in cells[:6])
            stt = values[0]
            if isinstance(stt, bool) or not isinstance(stt, (int, float)) or not float(stt).is_integer() or stt < 1:
                raise ValueError(f"Hàng {source_row}: STT phải là số nguyên dương.")
            stt = int(stt)
            if stt in seen_stt:
                raise ValueError(f"Hàng {source_row}: trùng STT {stt}.")
            seen_stt.add(stt)
            for column, value in zip(HEADERS[1:], values[1:]):
                if not isinstance(value, str) or not value.strip():
                    raise ValueError(f"Hàng {source_row}: thiếu hoặc sai kiểu trường {column!r}.")
            # Keep values as-is: strip is used only to reject empty cells above.
            records.append((stt, *values[1:], source_row))
        if not records:
            raise ValueError("Sheet không có nội dung.")
        source = {"fileName": path.name, "sha256": sha256(raw).hexdigest(), "sheet": sheet_name}
        return title, source, records
    finally:
        workbook.close()


def build_document(title, source, records):
    chapters = []
    section_count = 0
    for index, (stt, chapter_label, section_label, topic, details, tip, source_row) in enumerate(records, 1):
        if not chapters or chapters[-1]["label"] != chapter_label:
            if any(chapter["label"] == chapter_label for chapter in chapters):
                raise ValueError(f"Hàng {source_row}: chương {chapter_label!r} xuất hiện lại không liên tiếp.")
            chapters.append({"id": f"chapter-{len(chapters) + 1}", "label": chapter_label, "sections": []})
        sections = chapters[-1]["sections"]
        if not sections or sections[-1]["label"] != section_label:
            if any(section["label"] == section_label for section in sections):
                raise ValueError(f"Hàng {source_row}: phần {section_label!r} xuất hiện lại không liên tiếp.")
            section_count += 1
            sections.append({"id": f"section-{section_count}", "label": section_label, "rows": []})
        sections[-1]["rows"].append({
            "id": f"memory-{index}", "stt": stt, "topic": topic,
            "details": details, "tip": tip, "sourceRow": source_row,
        })
    return {"schemaVersion": 1, "title": title, "source": source, "chapters": chapters}


def verify_document(document, title, source, records):
    """Compare every packaged cell directly to the workbook, independently of building JSON."""
    if not isinstance(document, dict) or set(document) != {"schemaVersion", "title", "source", "chapters"}:
        raise ValueError("JSON sai cấu trúc cấp gốc.")
    if type(document["schemaVersion"]) is not int or document["schemaVersion"] != 1 or document["title"] != title or document["source"] != source:
        raise ValueError("JSON sai phiên bản, tiêu đề hoặc nguồn Excel/SHA-256.")
    flattened, chapter_count, section_count = [], 0, 0
    for chapter in document["chapters"]:
        chapter_count += 1
        if set(chapter) != {"id", "label", "sections"} or chapter["id"] != f"chapter-{chapter_count}" or not chapter["sections"]:
            raise ValueError("JSON sai cấu trúc, ID hoặc thứ tự chương.")
        for section in chapter["sections"]:
            section_count += 1
            if set(section) != {"id", "label", "rows"} or section["id"] != f"section-{section_count}" or not section["rows"]:
                raise ValueError("JSON sai cấu trúc, ID hoặc thứ tự phần.")
            for row in section["rows"]:
                if set(row) != {"id", "stt", "topic", "details", "tip", "sourceRow"} or row["id"] != f"memory-{len(flattened) + 1}":
                    raise ValueError("JSON sai cấu trúc, ID hoặc thứ tự nội dung.")
                if type(row["stt"]) is not int or type(row["sourceRow"]) is not int:
                    raise ValueError("JSON sai kiểu số nguyên STT hoặc hàng nguồn.")
                flattened.append((row["stt"], chapter["label"], section["label"], row["topic"], row["details"], row["tip"], row["sourceRow"]))
    if len(flattened) != len(records):
        raise ValueError(f"JSON có {len(flattened)} hàng; Excel có {len(records)} hàng.")
    for packaged, original in zip(flattened, records):
        if packaged != original:
            raise ValueError(f"Nội dung JSON khác Excel tại hàng {original[-1]} (STT {original[0]}).")
    expected_groups = [(record[1], record[2]) for index, record in enumerate(records)
                       if index == 0 or (record[1], record[2]) != (records[index - 1][1], records[index - 1][2])]
    actual_groups = [(chapter["label"], section["label"]) for chapter in document["chapters"] for section in chapter["sections"]]
    expected_chapters = [record[1] for index, record in enumerate(records) if index == 0 or record[1] != records[index - 1][1]]
    if actual_groups != expected_groups or [chapter["label"] for chapter in document["chapters"]] != expected_chapters:
        raise ValueError("JSON không giữ đúng cấu trúc chương/phần trong Excel.")
    return chapter_count, section_count, len(flattened)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="File .xlsx gốc, chỉ đọc")
    parser.add_argument("--output", type=Path, default=ROOT / "data/memory-tips.json")
    parser.add_argument("--sheet", default="Trang tính1")
    parser.add_argument("--check", action="store_true", help="Đối chiếu toàn bộ JSON hiện có với Excel, không ghi file")
    args = parser.parse_args()
    try:
        if args.source.resolve() == args.output.resolve():
            raise ValueError("File JSON đầu ra phải khác file Excel nguồn.")
        title, source, records = source_records(args.source, args.sheet)
        if args.check:
            document = json.loads(args.output.read_text(encoding="utf-8"))
        else:
            document = build_document(title, source, records)
        chapters, sections, rows = verify_document(document, title, source, records)
        if not args.check:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{'Khớp nguyên văn' if args.check else 'Đã xuất'}: {chapters} chương, {sections} phần, {rows} nội dung.")
        print(f"SHA-256 Excel: {source['sha256']}")
    except (OSError, ValueError, KeyError, TypeError) as error:
        parser.exit(1, f"Lỗi: {error}\n")


if __name__ == "__main__":
    main()
