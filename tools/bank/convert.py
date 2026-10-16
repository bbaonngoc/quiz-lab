#!/usr/bin/env python3
"""Local Admin only. Validate one workbook, preserve IDs, publish immutable JSON."""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import unicodedata
import uuid
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

ROOT = Path(__file__).resolve().parents[2]
ROLES = {
    "Bí thư Đảng ủy": "bi-thu-dang-uy",
    "Chủ tịch HĐND": "chu-tich-hdnd",
    "Chủ tịch UBND": "chu-tich-ubnd",
    "Chủ tịch Ủy ban MTTQ": "chu-tich-mttq",
}
FIELDS = {
    "id": ["ID"], "stt": ["STT"], "difficulty": ["Mức độ"],
    "scope": ["Phân loại"], "role": ["Chức danh", "Vai trò"],
    "question": ["Nội dung câu hỏi", "Nội dung", "Câu hỏi"],
    **{key: [f"Phương án {key}", key] for key in "ABCD"},
    "answer": ["Đáp án đúng"],
}


def key(value):
    text = unicodedata.normalize("NFD", str(value or "")).lower().replace("đ", "d")
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in text if not unicodedata.combining(c)))


def text(value):
    return "" if value is None else str(value).strip()


def encoded(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def digest(data):
    return hashlib.sha256(data).hexdigest()


def fingerprint(q):
    return json.dumps([q["question"], q["options"], q["role"]], ensure_ascii=False, sort_keys=True)


def previous_questions(output):
    manifest = output / "bank-manifest.json"
    if not manifest.exists():
        return []
    bank = json.loads(manifest.read_text(encoding="utf-8"))
    return [q for section in bank["sections"]
            for q in json.loads((output / section["url"]).read_text(encoding="utf-8"))["questions"]]


def convert(source, config, output, version, normalized, check=False):
    from openpyxl import load_workbook

    if source.suffix.lower() != '.xlsx' or normalized.suffix.lower() != '.xlsx':
        raise ValueError('Nguồn và bản chuẩn hóa phải là Excel .xlsx.')
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}", version):
        raise ValueError("Phiên bản chỉ gồm chữ, số, dấu chấm, gạch ngang/gạch dưới.")
    if source.resolve() == normalized.resolve():
        raise ValueError("Bản Excel có ID phải khác file nguồn; không ghi đè bản gốc.")
    if source.stat().st_size > 25 * 1024 * 1024:
        raise ValueError("Excel vượt quá 25 MB.")
    with ZipFile(source) as archive:
        if sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("Dữ liệu Excel giải nén vượt quá 64 MB.")

    chapters, sections = config["chapters"], config["sections"]
    chapter_ids = {c["id"] for c in chapters}
    if len(chapters) != 5 or len(chapter_ids) != 5 or len(sections) != 8:
        raise ValueError("Cấu hình phải có 5 chương và 8 đơn vị nội dung.")
    if len({s["id"] for s in sections}) != 8 or len({s["sheet"] for s in sections}) != 8:
        raise ValueError("ID đơn vị nội dung và tên sheet phải duy nhất.")
    if any(s["chapterId"] not in chapter_ids for s in sections):
        raise ValueError("Đơn vị nội dung thuộc chương không hợp lệ.")

    old = previous_questions(output)
    old_fingerprints = defaultdict(list)
    for q in old:
        old_fingerprints[fingerprint(q)].append(q["id"])
    report = {"source": source.name, "sourceSha256": digest(source.read_bytes()),
              "dataVersion": version, "errors": [], "warnings": [], "sections": []}
    aliases = {key(label): field for field, labels in FIELDS.items() for label in labels}
    role_aliases = {key(label): role for label, role in ROLES.items()}
    role_aliases.update({key(role): role for role in ROLES.values()})
    common = {key("Chung"), key("Dùng chung"), "all"}
    levels = {key(level): level for level in config["difficulties"]}
    questions, seen_ids, seen_content = [], {}, {}
    workbook = load_workbook(source, data_only=False, keep_links=False)
    try:
        expected = {s["sheet"] for s in sections}
        missing = expected - set(workbook.sheetnames)
        extra = set(workbook.sheetnames) - expected - set(config.get("ignoredSheets", []))
        if missing or extra:
            raise ValueError(f"Sheet thiếu: {sorted(missing)}; sheet ngoài cấu hình: {sorted(extra)}.")
        for section in sections:
            sheet = workbook[section["sheet"]]
            columns = {}
            for cell in sheet[1]:
                field = aliases.get(key(cell.value))
                if field:
                    if field in columns:
                        report["errors"].append({"sheet": sheet.title, "row": 1, "column": cell.column_letter, "message": f"Cột {field} bị trùng."})
                    columns[field] = cell.column
            for field in ["question", "A", "B", "C", "D", "answer"]:
                if field not in columns:
                    report["errors"].append({"sheet": sheet.title, "row": 1, "column": field, "message": "Thiếu cột bắt buộc."})
            had_id = "id" in columns
            if not had_id:
                columns["id"] = sheet.max_column + 1
                sheet.cell(1, columns["id"], "ID")
            section_questions, stt_seen = [], set()
            for rn in range(2, sheet.max_row + 1):
                values = {field: text(sheet.cell(rn, col).value) for field, col in columns.items()}
                if not any(values.values()):
                    continue
                location = {"sheet": sheet.title, "row": rn}
                row_errors = []

                def fail(field, message):
                    col = sheet.cell(rn, columns[field]).column_letter if field in columns else field
                    row_errors.append({**location, "column": col, "message": message})

                for field in ["question", "A", "B", "C", "D", "answer"]:
                    if not values.get(field):
                        fail(field, "Thiếu nội dung bắt buộc.")
                for field, col in columns.items():
                    if sheet.cell(rn, col).data_type in ("f", "e"):
                        fail(field, "Ô công thức/lỗi Excel: hãy thay bằng giá trị trước khi phát hành.")
                answer = values.get("answer", "").upper()
                if answer not in list("ABCD"):
                    fail("answer", "Đáp án đúng phải là A, B, C hoặc D.")
                difficulty = levels.get(key(values.get("difficulty"))) if values.get("difficulty") else None
                if values.get("difficulty") and difficulty is None:
                    fail("difficulty", "Mức độ không thuộc danh mục cấu hình.")
                scope, raw_role = key(values.get("scope")), key(values.get("role"))
                role = role_aliases.get(raw_role)
                if not raw_role:
                    fail("role", "Thiếu Chức danh (role). Câu dùng chung phải ghi rõ Chung; không tự mặc định.")
                if not scope:
                    scope = "chucdanh" if role else "dungchung" if raw_role in common else ""
                    report["warnings"].append({**location, "message": "Thiếu Phân loại; kiểm tra quy tắc suy ra phạm vi."})
                if scope == "dungchung":
                    if raw_role and raw_role not in common:
                        fail("role", "Câu dùng chung không được gắn chức danh riêng.")
                    role, scope_value = "all", "shared"
                elif scope == "chucdanh":
                    if not role:
                        fail("role", "Câu riêng phải có một trong bốn chức danh hợp lệ.")
                    scope_value = "role"
                else:
                    fail("scope", "Chưa xác định được Phân loại Dùng chung/Chức danh.")
                    scope_value = None
                stt = None
                if values.get("stt"):
                    try:
                        numeric = float(values["stt"])
                        if not numeric.is_integer() or numeric <= 0:
                            raise ValueError()
                        stt = int(numeric)
                        if stt in stt_seen:
                            fail("stt", "STT trùng trong cùng sheet.")
                        stt_seen.add(stt)
                    except ValueError:
                        fail("stt", "STT phải là số nguyên dương hoặc để trống.")
                q = {"chapterId": section["chapterId"], "sectionId": section["id"], "stt": stt,
                     "difficulty": difficulty, "scope": scope_value, "role": role,
                     "topic": section["title"], "question": values.get("question", ""),
                     "options": {k: values.get(k, "") for k in "ABCD"}, "answer": answer}
                qid = values.get("id")
                if not qid:
                    matches = old_fingerprints.get(fingerprint(q), [])
                    if len(matches) == 1:
                        qid = matches[0]
                    elif old and not had_id:
                        fail("id", "Không thể đối chiếu ID. Hãy sửa trên bản Excel đã được cấp ID.")
                    else:
                        qid = f"q-{uuid.uuid4()}"
                if qid and (not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._:-]{0,99}", qid) or qid in {"prototype", "constructor", "toString", "__proto__"}):
                    fail("id", "ID không hợp lệ.")
                if qid in seen_ids:
                    fail("id", f"ID trùng với {seen_ids[qid]}.")
                if row_errors:
                    report["errors"].extend(row_errors)
                    continue
                seen_ids[qid] = f"{sheet.title}, dòng {rn}"
                revision_fields = {k: q[k] for k in ["question", "options", "answer", "scope", "role"]}
                q = {"id": qid, **q, "revision": digest(encoded(revision_fields))}
                content = key(q["question"]) + json.dumps([key(q["options"][k]) for k in "ABCD"])
                if content in seen_content:
                    report["warnings"].append({**location, "message": f"Nội dung trùng với {seen_content[content]}; không tự xóa."})
                seen_content[content] = f"{sheet.title}, dòng {rn}"
                sheet.cell(rn, columns["id"], qid)
                section_questions.append(q)
            if not section_questions:
                report["errors"].append({"sheet": sheet.title, "row": 1, "message": "Không có câu hợp lệ."})
            questions.extend(section_questions)
            report["sections"].append({"id": section["id"], "count": len(section_questions),
                                       "roles": dict(Counter(q["role"] for q in section_questions))})
        if len(questions) > 10000:
            report["errors"].append({"message": "Ngân hàng vượt quá 10.000 câu."})
        report["count"] = len(questions)
        old_by_id, new_by_id = {q["id"]: q for q in old}, {q["id"]: q for q in questions}
        report["changes"] = {
            "added": sorted(new_by_id.keys() - old_by_id.keys()),
            "removed": sorted(old_by_id.keys() - new_by_id.keys()),
            "modified": [{"id": qid, "answerChanged": old_by_id[qid]["answer"] != q["answer"],
                          "scopeChanged": old_by_id[qid]["role"] != q["role"]}
                         for qid, q in new_by_id.items() if qid in old_by_id and q != old_by_id[qid]],
        }
        if report["errors"] or check:
            return report
        output.mkdir(parents=True, exist_ok=True)
        normalized.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="quizlab-bank-", dir=output.parent) as temp:
            stage = Path(temp)
            release = stage / "release"
            release.mkdir()
            manifest_sections = []
            for section in sections:
                subset = [q for q in questions if q["sectionId"] == section["id"]]
                payload = encoded({"schemaVersion": 2, "dataVersion": version, "sectionId": section["id"], "questions": subset})
                name = section["id"] + ".json"
                (release / name).write_bytes(payload)
                manifest_sections.append({k: v for k, v in section.items() if k != "sheet"} | {
                    "url": f"releases/{version}/{name}", "sha256": digest(payload), "count": len(subset)})
            manifest = {"schemaVersion": 2, "dataVersion": version, "title": config["title"],
                        "difficulties": config["difficulties"], "count": len(questions),
                        "chapters": chapters, "sections": manifest_sections}
            (stage / "bank-manifest.json").write_bytes(encoded(manifest))
            workbook.save(stage / "normalized.xlsx")
            destination = output / "releases" / version
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                if {p.name: p.read_bytes() for p in destination.iterdir()} != {p.name: p.read_bytes() for p in release.iterdir()}:
                    raise ValueError("Phiên bản đã tồn tại với nội dung khác; hãy tăng --version.")
            else:
                os.replace(release, destination)
            shutil.copyfile(stage / "normalized.xlsx", normalized)
            os.replace(stage / "bank-manifest.json", output / "bank-manifest.json")
        return report
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("config.json"))
    parser.add_argument("--output", type=Path, default=ROOT / "tools/bank/output/data")
    parser.add_argument("--version", required=True)
    parser.add_argument("--normalized", type=Path, default=ROOT / "tools/bank/output/cau-hoi-co-id.xlsx")
    parser.add_argument("--report", type=Path, default=ROOT / "tools/bank/output/report.json")
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    try:
        report = convert(args.workbook, json.loads(args.config.read_text(encoding="utf-8")),
                         args.output, args.version, args.normalized, args.check)
    except (OSError, ValueError, KeyError, ImportError, BadZipFile, ParseError) as error:
        report = {"errors": [{"message": str(error)}]}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_bytes(encoded(report))
    for issue in report.get("errors", []):
        print(f"LỖI {issue.get('sheet', '')} dòng {issue.get('row', '')} cột {issue.get('column', '')}: {issue['message']}", file=sys.stderr)
    print(f"{len(report.get('errors', []))} lỗi; {len(report.get('warnings', []))} cảnh báo. Báo cáo: {args.report}")
    if report.get("errors"):
        return 1
    print(f"{report['count']} câu / 8 mục / 5 chương. " + ("Chỉ kiểm tra, chưa phát hành." if args.check else f"Đầu ra: {args.output}; Excel có ID: {args.normalized}"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
