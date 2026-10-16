import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from convert import convert

CONFIG = json.loads(Path(__file__).with_name('config.json').read_text())
HEADERS = ['STT', 'Mức độ', 'Phân loại', 'Chức danh', 'Nội dung câu hỏi', 'Phương án A', 'Phương án B', 'Phương án C', 'Phương án D', 'Đáp án đúng']


class ConverterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.xlsx'
        self.output = self.root / 'data'
        self.normalized = self.root / 'normalized.xlsx'
        book = Workbook()
        book.remove(book.active)
        for i, section in enumerate(CONFIG['sections']):
            sheet = book.create_sheet(section['sheet'])
            sheet.append(HEADERS)
            sheet.append([1, None, 'Dùng chung', 'Chung', f'Câu hỏi {i}?', 'a', 'b', 'c', 'd', ' b '])
        book.save(self.source)
        book.close()

    def tearDown(self):
        self.temp.cleanup()

    def run_convert(self, source=None, version='test-1', normalized=None, check=False):
        return convert(source or self.source, copy.deepcopy(CONFIG), self.output, version,
                       normalized or self.normalized, check)

    def edit(self, path, callback):
        book = load_workbook(path)
        callback(book)
        book.save(path)
        book.close()

    def test_ids_and_immutable_releases(self):
        report = self.run_convert()
        self.assertFalse(report['errors'])
        self.assertEqual(report['count'], 8)
        first = (self.output / 'bank-manifest.json').read_bytes()
        self.assertTrue(self.normalized.exists())
        # Reimporting the same original source reuses verified IDs, byte for byte.
        second = self.run_convert()
        self.assertEqual(second['changes']['added'], [])
        self.assertEqual(first, (self.output / 'bank-manifest.json').read_bytes())
        self.edit(self.normalized, lambda b: setattr(b.worksheets[0]['E2'], 'value', 'Nội dung đã sửa?'))
        with self.assertRaisesRegex(ValueError, 'Phiên bản đã tồn tại'):
            self.run_convert(self.normalized, normalized=self.root / 'next.xlsx')
        updated = self.run_convert(self.normalized, version='test-2', normalized=self.root / 'next.xlsx')
        self.assertFalse(updated['errors'])
        self.assertEqual(len(updated['changes']['modified']), 1)
        self.assertEqual(updated['changes']['added'], [])
        self.assertEqual(updated['changes']['removed'], [])

    def test_errors_do_not_replace_published_bank(self):
        self.run_convert()
        before = (self.output / 'bank-manifest.json').read_bytes()
        self.edit(self.source, lambda b: setattr(b.worksheets[0]['F2'], 'value', None))
        report = self.run_convert(version='bad')
        self.assertTrue(any(e.get('column') == 'F' and e['row'] == 2 for e in report['errors']))
        self.assertEqual(before, (self.output / 'bank-manifest.json').read_bytes())
        self.assertFalse((self.output / 'releases/bad').exists())

    def test_invalid_scope_role_answer_formula_and_duplicate_id(self):
        self.run_convert()
        def change(b):
            b.worksheets[0]['C2'] = 'Chức danh'
            b.worksheets[0]['D2'] = 'Chức danh không có'
            b.worksheets[1]['J2'] = 'E'
            b.worksheets[2]['E2'] = '=1+1'
            b.worksheets[4]['K2'] = b.worksheets[3]['K2'].value
        self.edit(self.normalized, change)
        report = self.run_convert(self.normalized, 'bad', self.root / 'next.xlsx')
        messages = ' '.join(e['message'] for e in report['errors'])
        for expected in ['chức danh hợp lệ', 'A, B, C hoặc D', 'công thức', 'ID trùng']:
            self.assertIn(expected, messages)

    def test_sheet_order_is_irrelevant_but_missing_sheet_fails(self):
        self.run_convert()
        self.edit(self.normalized, lambda b: b.move_sheet(b.worksheets[-1], offset=-7))
        report = self.run_convert(self.normalized, 'test-2', self.root / 'next.xlsx')
        self.assertEqual(report['changes']['modified'], [])
        self.edit(self.normalized, lambda b: b.remove(b.worksheets[0]))
        with self.assertRaisesRegex(ValueError, 'Sheet thiếu'):
            self.run_convert(self.normalized, 'test-3', self.root / 'next.xlsx')

    def test_required_columns_and_check_only(self):
        report = self.run_convert(check=True)
        self.assertFalse(report['errors'])
        self.assertFalse((self.output / 'bank-manifest.json').exists())
        self.edit(self.source, lambda b: setattr(b.worksheets[0]['J1'], 'value', 'Không phải đáp án'))
        report = self.run_convert(check=True)
        self.assertTrue(any(e.get('row') == 1 for e in report['errors']))

    def test_missing_role_is_reported_even_for_shared_questions(self):
        self.run_convert()
        before = (self.output / 'bank-manifest.json').read_bytes()
        for missing in [None, '', '   ']:
            self.edit(self.source, lambda b: setattr(b.worksheets[0]['D2'], 'value', missing))
            report = self.run_convert(version='missing-role')
            self.assertTrue(any(e.get('column') == 'D' and 'Thiếu Chức danh' in e['message'] for e in report['errors']))
            self.assertEqual(before, (self.output / 'bank-manifest.json').read_bytes())

    def test_role_scope_and_difficulty_preserved_from_excel(self):
        pairs = [('Bí thư Đảng ủy', 'bi-thu-dang-uy'), ('Chủ tịch HĐND', 'chu-tich-hdnd'),
                 ('Chủ tịch UBND', 'chu-tich-ubnd'), ('Chủ tịch Ủy ban MTTQ', 'chu-tich-mttq')]
        def append_rows(book):
            for sheet in book:
                sheet['B2'] = 'Mức 1'
                for index, (label, _) in enumerate(pairs, 2):
                    sheet.append([index, 'Mức 2', 'Chức danh', label, f'Câu {sheet.title} {index}?', 'a', 'b', 'c', 'd', 'A'])
        self.edit(self.source, append_rows)
        report = self.run_convert()
        self.assertFalse(report['errors'])
        manifest = json.loads((self.output / 'bank-manifest.json').read_text())
        for section in manifest['sections']:
            questions = json.loads((self.output / section['url']).read_text())['questions']
            self.assertEqual(questions[0]['role'], 'all')
            self.assertEqual(questions[0]['scope'], 'shared')
            self.assertEqual(questions[0]['difficulty'], 'Mức 1')
            self.assertEqual([q['role'] for q in questions[1:]], [role for _, role in pairs])
            self.assertTrue(all(q['scope'] == 'role' and q['difficulty'] == 'Mức 2' for q in questions[1:]))


if __name__ == '__main__':
    unittest.main()
